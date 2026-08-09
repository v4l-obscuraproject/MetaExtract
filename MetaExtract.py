"""
MetaExtract
Public File Metadata Extraction Tool

Created by: 5h9q_ (developer) and rosp_1 (publishing) on discord.
Licensed under the MIT License (see LICENSE file)

Extracts hidden metadata from public files -- PDFs, images, and Office
documents. Files often carry metadata their creator didn't realize was
there: author names, software used, GPS coordinates on photos, last
modified dates, and more. This is a standard OSINT/forensics technique
for analyzing publicly available documents (e.g. files found via
DorkRecon's "Publicly indexed documents" dork category).

IMPORTANT: Only run this against files you own or have explicit
authorization to analyze, or files that are genuinely and legitimately
public (e.g. found via authorized recon on a domain you're testing).
Extracting metadata from someone's file doesn't require their
permission technically, but using it to identify or track a specific
person without authorization is not an intended or acceptable use of
this tool.

Usage:
    python meta_extract.py <file_path>

Supports: .pdf, .jpg/.jpeg, .png, .docx, .xlsx
"""

import argparse
import os
import sys

BANNER = r"""
  ____  ____   _____  _____ _    _ _____
 / __ \|  _ \ / ____|/ ____| |  | |  __ \     /\
| |  | | |_) | (___ | |    | |  | | |__) |   /  \
| |  | |  _ < \___ \| |    | |  | |  _  /   / /\ \
| |__| | |_) |____) | |____| |__| | | \ \  / ____ \
 \____/|____/|_____/ \_____|\____/|_|  \_\/_/    \_\

 __  __     _        ___     _               _
|  \/  |___| |_ __ _| __|_ _| |_ _ _ __ _ __| |_
| |\/| / -_)  _/ _` | _|\ \ /  _| '_/ _` / _|  _|
|_|  |_\___|\__\__,_|___/_\_\\__|_| \__,_\__|\__|

    5h9q_  (developer)         
"""


def extract_pdf_metadata(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    meta = reader.metadata
    result = {}
    if meta:
        for key, value in meta.items():
            clean_key = key.lstrip("/")
            result[clean_key] = str(value)
    result["Page count"] = len(reader.pages)
    return result


def extract_image_metadata(path):
    from PIL import Image
    from PIL.ExifTags import TAGS, GPSTAGS

    result = {}
    img = Image.open(path)
    result["Format"] = img.format
    result["Size"] = f"{img.width}x{img.height}"
    result["Mode"] = img.mode

    exif_data = img.getexif()
    if not exif_data:
        return result

    for tag_id, value in exif_data.items():
        tag = TAGS.get(tag_id, tag_id)
        if tag == "GPSInfo":
            gps_data = {}
            for gps_id, gps_value in value.items():
                gps_tag = GPSTAGS.get(gps_id, gps_id)
                gps_data[gps_tag] = gps_value
            if gps_data:
                result["GPS"] = gps_data
        else:
            result[str(tag)] = str(value)

    return result


def extract_docx_metadata(path):
    from docx import Document
    doc = Document(path)
    props = doc.core_properties
    return {
        "Author": props.author,
        "Last modified by": props.last_modified_by,
        "Created": str(props.created) if props.created else None,
        "Modified": str(props.modified) if props.modified else None,
        "Title": props.title,
        "Subject": props.subject,
        "Company": getattr(props, "company", None),
        "Revision": props.revision,
    }


def extract_xlsx_metadata(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    props = wb.properties
    return {
        "Author": props.creator,
        "Last modified by": props.lastModifiedBy,
        "Created": str(props.created) if props.created else None,
        "Modified": str(props.modified) if props.modified else None,
        "Title": props.title,
        "Subject": props.subject,
        "Company": getattr(props, "company", None),
    }


EXTRACTORS = {
    ".pdf": extract_pdf_metadata,
    ".jpg": extract_image_metadata,
    ".jpeg": extract_image_metadata,
    ".png": extract_image_metadata,
    ".docx": extract_docx_metadata,
    ".xlsx": extract_xlsx_metadata,
}


def main():
    parser = argparse.ArgumentParser(
        description="Extract hidden metadata from public files (PDF, image, Office docs)"
    )
    parser.add_argument("file_path", help="Path to the file to analyze")
    args = parser.parse_args()

    print(BANNER)

    if not os.path.isfile(args.file_path):
        print(f"[!] File not found: {args.file_path}")
        sys.exit(1)

    ext = os.path.splitext(args.file_path)[1].lower()
    extractor = EXTRACTORS.get(ext)

    if not extractor:
        print(f"[!] Unsupported file type: {ext}")
        print(f"[!] Supported types: {', '.join(EXTRACTORS.keys())}")
        sys.exit(1)

    print(f"[*] Analyzing: {args.file_path}\n")

    try:
        metadata = extractor(args.file_path)
    except Exception as e:
        print(f"[!] Failed to extract metadata: {e}")
        sys.exit(1)

    found_any = False
    for key, value in metadata.items():
        if value not in (None, "", "None"):
            print(f"  {key}: {value}")
            found_any = True

    if not found_any:
        print("  No metadata found (file may have been stripped/sanitized).")

    print()


if __name__ == "__main__":
    main()